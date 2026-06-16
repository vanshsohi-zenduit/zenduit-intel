"""Generate Outbound Intelligence Engine — executive slide deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY    = RGBColor(0x0F, 0x17, 0x2A)
BLUE    = RGBColor(0x25, 0x63, 0xEB)
LBLUE   = RGBColor(0x93, 0xC5, 0xFD)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SLATE   = RGBColor(0x94, 0xA3, 0xB8)
LIGHT   = RGBColor(0xDB, 0xEA, 0xFE)
GREEN   = RGBColor(0x22, 0xC5, 0x5E)
YELLOW  = RGBColor(0xFB, 0xBF, 0x24)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]  # completely blank


def add_slide():
    return prs.slides.add_slide(BLANK)


def bg(slide, color=NAVY):
    from pptx.util import Emu
    shape = slide.shapes.add_shape(1, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def box(slide, l, t, w, h, fill=BLUE, alpha=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def txt(slide, text, l, t, w, h, size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    p  = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.name  = "Calibri"
    return tb


def label(slide, text, l, t, color=SLATE, size=11):
    txt(slide, text.upper(), l, t, 8, 0.3, size=size, color=color, bold=False)


def divider(slide, t, color=BLUE, thickness=0.03):
    box(slide, 0.6, t, 12.13, thickness, fill=color)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide()
bg(s)

# left accent bar
box(s, 0, 0, 0.08, 7.5, fill=BLUE)

# top-right glow blob (decorative rectangle)
b = box(s, 9.5, -0.5, 4.5, 4.5, fill=RGBColor(0x1D, 0x4E, 0xD8))
b.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

txt(s, "ZENDUIT  ·  AI & AUTOMATION", 1.1, 1.6, 10, 0.5, size=11, color=SLATE)
txt(s, "Outbound Intelligence\nEngine", 1.1, 2.1, 10, 2.2, size=48, bold=True, color=WHITE)
txt(s, "Automated prospect research, personalised outreach &\nreal-time sales coaching — in under 3 minutes.", 1.1, 4.4, 9, 1.0, size=18, color=LBLUE)
txt(s, "Sales AI  ·  June 2026", 1.1, 6.6, 5, 0.4, size=12, color=SLATE)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — The Problem
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 0.08, 7.5, fill=BLUE)

label(s, "The Problem", 1.1, 0.5)
txt(s, "Manual research was killing\nsales productivity.", 1.1, 0.9, 11, 1.6, size=38, bold=True, color=WHITE)

stats = [
    ("2–3 hrs", "per prospect for research,\nlinkedin & email crafting"),
    ("Generic", "outreach with no\npersonalised signals"),
    ("No playbook", "inconsistent messaging\nacross reps"),
]
for i, (val, desc) in enumerate(stats):
    x = 1.1 + i * 4.0
    box(s, x, 3.0, 3.6, 2.8, fill=RGBColor(0x1E, 0x29, 0x3B))
    txt(s, val,  x+0.2, 3.2,  3.2, 0.8, size=30, bold=True, color=YELLOW)
    txt(s, desc, x+0.2, 4.0,  3.2, 1.4, size=14, color=SLATE)

txt(s, "Every rep was starting from scratch — no consistent intel, no leverage.", 1.1, 6.3, 11, 0.6, size=13, color=SLATE)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — The Solution
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 0.08, 7.5, fill=BLUE)

label(s, "The Solution", 1.1, 0.5)
txt(s, "A 4-phase AI pipeline that turns a\nwebsite URL into a full sales brief.", 1.1, 0.9, 11, 1.5, size=34, bold=True, color=WHITE)

phases = [
    (BLUE,                       "01", "Website Research",     "News, hiring signals,\ncompetitor intel"),
    (RGBColor(0x70, 0x5A, 0xD5), "02", "LinkedIn Intel",       "Decision-maker contacts,\npersonalisation hooks"),
    (RGBColor(0x0D, 0x9A, 0x75), "03", "Product Intelligence", "Brain MCP: similar\naccounts & battlecards"),
    (RGBColor(0xD9, 0x77, 0x06), "04", "Strategy Generation",  "Briefing, scripts,\n14-day sequence"),
]
for i, (col, num, title, desc) in enumerate(phases):
    x = 0.6 + i * 3.2
    box(s, x, 2.7, 3.0, 3.5, fill=RGBColor(0x1E, 0x29, 0x3B))
    box(s, x, 2.7, 3.0, 0.06, fill=col)
    txt(s, num,   x+0.2, 2.85, 2.6, 0.5, size=22, bold=True, color=col)
    txt(s, title, x+0.2, 3.35, 2.6, 0.6, size=14, bold=True, color=WHITE)
    txt(s, desc,  x+0.2, 4.0,  2.6, 1.8, size=12, color=SLATE)

txt(s, "< 3 minutes  ·  fully automated  ·  runs on demand", 1.1, 6.5, 11, 0.5,
    size=13, color=LBLUE, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — What It Produces
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 0.08, 7.5, fill=BLUE)

label(s, "Output", 1.1, 0.5)
txt(s, "Everything a rep needs before\nthe first touchpoint.", 1.1, 0.9, 11, 1.5, size=34, bold=True, color=WHITE)

outputs = [
    ("Executive Briefing",   "Company overview, pain points, competitive\nlandscape — tailored to Zenduit's products."),
    ("5 Outreach Scripts",   "LinkedIn, cold email & call scripts with named\ncontacts, hooks from real posts & ROI data."),
    ("14-Day Sequence",      "Day-by-day outreach plan with channel,\nsubject lines and personalised instructions."),
    ("Objection Playbook",   "Pre-built objection table mapped to the\nprospect's specific vendors and pain points."),
    ("Live Call Coach",      "Real-time AI coaching overlay during calls\n— listens, identifies speaker, suggests next move."),
    ("Prospect Library",     "All runs saved and searchable — reps can\nreload any prospect in one click."),
]
for i, (title, desc) in enumerate(outputs):
    row, col = divmod(i, 2)
    x = 1.1 + col * 6.2
    y = 2.7 + row * 1.55
    box(s, x, y, 5.9, 1.35, fill=RGBColor(0x1E, 0x29, 0x3B))
    txt(s, title, x+0.2, y+0.12, 5.5, 0.4, size=13, bold=True, color=WHITE)
    txt(s, desc,  x+0.2, y+0.55, 5.5, 0.7, size=11, color=SLATE)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Adoption & Integration
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 0.08, 7.5, fill=BLUE)

label(s, "Adoption & Integration", 1.1, 0.5)
txt(s, "Built for the sales team,\nintegrated into their workflow.", 1.1, 0.9, 11, 1.5, size=34, bold=True, color=WHITE)

items = [
    (GREEN,  "Live & Deployed",      "Web app accessible to sales reps on demand. No CLI, no setup — enter a URL, get a brief."),
    (BLUE,   "LinkedIn Connected",   "Direct LinkedIn MCP integration pulls real decision-maker profiles and recent post hooks."),
    (BLUE,   "Brain MCP Connected",  "Queries internal Zenduit CRM data — similar won accounts, battlecards, objection scripts."),
    (YELLOW, "In Active Testing",    "Sales reps running it on live prospects. Feedback loop in place with the tribe leader."),
]
for i, (col, title, desc) in enumerate(items):
    y = 2.8 + i * 1.05
    box(s, 1.1, y, 0.06, 0.75, fill=col)
    txt(s, title, 1.4, y+0.02, 4.5, 0.35, size=13, bold=True, color=WHITE)
    txt(s, desc,  1.4, y+0.40, 10.5, 0.55, size=12, color=SLATE)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Business Impact
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 0.08, 7.5, fill=BLUE)

label(s, "Business Impact", 1.1, 0.5)
txt(s, "Faster pipeline, better signals,\nconsistent execution.", 1.1, 0.9, 11, 1.5, size=34, bold=True, color=WHITE)

metrics = [
    (GREEN,  "~2 hrs saved",  "per prospect researched\nby each rep"),
    (BLUE,   "Personalised",  "every outreach script\nuses real LinkedIn data"),
    (YELLOW, "Standardised",  "playbook used across\nall outbound reps"),
    (LBLUE,  "Scalable",      "bulk mode processes\n20 prospects at once"),
]
for i, (col, val, desc) in enumerate(metrics):
    x = 1.1 + i * 3.0
    box(s, x, 3.0, 2.7, 2.6, fill=RGBColor(0x1E, 0x29, 0x3B))
    box(s, x, 3.0, 2.7, 0.06, fill=col)
    txt(s, val,  x+0.2, 3.2,  2.3, 0.7, size=22, bold=True, color=col)
    txt(s, desc, x+0.2, 3.95, 2.3, 1.4, size=12, color=SLATE)

txt(s, "Direct impact: reps spend time selling, not researching. Manager visibility into rep prep quality.", 1.1, 6.2, 11.2, 0.6, size=12, color=SLATE)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Next Steps & Blockers
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide()
bg(s)
box(s, 0, 0, 0.08, 7.5, fill=BLUE)

label(s, "Next Steps & Blockers", 1.1, 0.5)
txt(s, "What's next and\nwhat we need.", 1.1, 0.9, 9, 1.3, size=34, bold=True, color=WHITE)

next_steps = [
    "Formalise adoption — onboard remaining outbound reps with a walkthrough session",
    "Live Coach refinement — improve speaker detection accuracy and response timing",
    "CRM write-back — auto-log generated briefs and sequences into Zoho CRM",
    "Metrics tracking — instrument usage to measure reps using it vs. not and deal outcomes",
]
blockers = [
    "LinkedIn rate limits — need a stable scraping solution for high-volume use",
    "Rep buy-in — a short training session needed to drive consistent adoption",
]

txt(s, "NEXT STEPS", 1.1, 2.5, 6, 0.35, size=10, color=SLATE, bold=True)
for i, item in enumerate(next_steps):
    box(s, 1.1, 3.0 + i*0.82, 0.06, 0.55, fill=GREEN)
    txt(s, item, 1.4, 3.0 + i*0.82, 5.8, 0.7, size=12, color=WHITE)

txt(s, "BLOCKERS / GAPS", 7.4, 2.5, 5.5, 0.35, size=10, color=SLATE, bold=True)
for i, item in enumerate(blockers):
    box(s, 7.4, 3.0 + i*0.92, 0.06, 0.65, fill=YELLOW)
    txt(s, item, 7.7, 3.0 + i*0.92, 5.3, 0.8, size=12, color=WHITE)

divider(s, 6.8, color=RGBColor(0x1E, 0x29, 0x3B), thickness=0.02)
txt(s, "Goal: every rep running at least one prospect through the tool per day by end of June.",
    1.1, 6.9, 11.2, 0.4, size=12, color=LBLUE)


# ═══════════════════════════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════════════════════════
out = "Outbound_Intelligence_Engine.pptx"
prs.save(out)
print(f"Saved: {out}")
